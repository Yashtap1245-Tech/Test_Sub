
import openpyxl

wb = openpyxl.load_workbook(r'd:\Test_Sub\Vendor-Payments_2021-23.xlsx', read_only=True, data_only=True)

for sh in wb.sheetnames:
    ws = wb[sh]
    agencies = set()
    categories = set()
    subcats = set()
    months = set()
    total_amount = 0.0
    row_count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        agencies.add(str(row[4]).strip())
        categories.add(str(row[6]).strip())
        subcats.add(str(row[8]).strip())
        months.add(str(row[2]).strip())
        try:
            total_amount += float(row[10])
        except:
            pass
        row_count += 1

    print("=== " + sh + " ===")
    print("Rows:", row_count)
    print("Total Amount: $" + "{:,.2f}".format(total_amount))
    print("Unique Agencies (" + str(len(agencies)) + "): " + str(list(agencies)[:8]))
    print("Unique Categories (" + str(len(categories)) + "): " + str(list(categories)))
    print("Unique SubCategories (" + str(len(subcats)) + "): " + str(list(subcats)))
    print("Fiscal Months: " + str(sorted(months)))
    print()
