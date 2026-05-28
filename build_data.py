import openpyxl
import json
import os
from collections import defaultdict

def build_datasets():
    xlsx_path = r"d:\Test_Sub\Vendor-Payments_2021-23.xlsx"
    output_dir = r"d:\Test_Sub\app\data"
    os.makedirs(output_dir, exist_ok=True)

    print("Loading workbook...")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    
    # Aggregation dictionaries
    # 1. Summary: (agency, category, subcategory, fy, month) -> amount
    summary_agg = defaultdict(float)
    # 2. Vendors: (agency, category, vendor, fy) -> amount
    vendor_agg = defaultdict(float)
    
    # Track unique metadata values
    agencies_set = set()
    categories_set = set()
    subcategories_set = set()
    vendors_set = set()
    
    total_processed_rows = 0

    for sheet_name in wb.sheetnames:
        print(f"Processing sheet: {sheet_name}...")
        ws = wb[sheet_name]
        
        # Read header to make sure columns match
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter)
        print("Header:", header)
        
        # Column mappings
        # ('Bien', 'FY', 'FMonth', 'Agy', 'Agency', 'Object', 'Category', 'Subobj', 'SubCategory', 'Vendor', 'Amount')
        # We index based on header names to be safe
        col_indices = {name: i for i, name in enumerate(header)}
        
        fy_idx = col_indices['FY']
        month_idx = col_indices['FMonth']
        agency_idx = col_indices['Agency']
        cat_idx = col_indices['Category']
        subcat_idx = col_indices['SubCategory']
        vendor_idx = col_indices['Vendor']
        amount_idx = col_indices['Amount']

        count = 0
        for row in rows_iter:
            if row[0] is None:
                continue
                
            fy = str(row[fy_idx]).strip()
            month = str(row[month_idx]).strip()
            agency = str(row[agency_idx]).strip()
            category = str(row[cat_idx]).strip()
            subcategory = str(row[subcat_idx]).strip()
            vendor = str(row[vendor_idx]).strip()
            
            try:
                amount = float(row[amount_idx])
            except (ValueError, TypeError):
                amount = 0.0
                
            # Add to sets for metadata
            agencies_set.add(agency)
            categories_set.add(category)
            subcategories_set.add(subcategory)
            vendors_set.add(vendor)
            
            # Aggregate summary data
            summary_key = (agency, category, subcategory, fy, month)
            summary_agg[summary_key] += amount
            
            # Aggregate vendor data
            vendor_key = (agency, category, vendor, fy)
            vendor_agg[vendor_key] += amount
            
            count += 1
            if count % 100000 == 0:
                print(f"  Processed {count} rows...")
                
        total_processed_rows += count
        print(f"Finished {sheet_name}. Processed rows: {count}")

    print("\nProcessing complete. Preparing exports...")
    
    # 1. Format metadata
    metadata = {
        "agencies": sorted(list(agencies_set)),
        "categories": sorted(list(categories_set)),
        "subcategories": sorted(list(subcategories_set)),
        "fiscal_years": ["2022", "2023"]
    }
    
    # 2. Format summary list
    summary_list = []
    for (agency, category, subcategory, fy, month), amount in summary_agg.items():
        if amount == 0.0:
            continue
        summary_list.append({
            "a": agency,
            "c": category,
            "s": subcategory,
            "y": fy,
            "m": month,
            "v": round(amount, 2)
        })
        
    # 3. Format vendor list
    vendor_list = []
    for (agency, category, vendor, fy), amount in vendor_agg.items():
        if amount == 0.0:
            continue
        vendor_list.append({
            "a": agency,
            "c": category,
            "v": vendor,
            "y": fy,
            "val": round(amount, 2)
        })
        
    # Write to files
    meta_file = os.path.join(output_dir, "meta.json")
    summary_file = os.path.join(output_dir, "summary.json")
    vendors_file = os.path.join(output_dir, "vendors.json")
    
    print(f"Writing {meta_file}...")
    with open(meta_file, "w", encoding="utf-8") as f:
        json.dump(metadata, f, ensure_ascii=False)
        
    print(f"Writing {summary_file}...")
    with open(summary_file, "w", encoding="utf-8") as f:
        json.dump(summary_list, f, ensure_ascii=False)
        
    print(f"Writing {vendors_file}...")
    with open(vendors_file, "w", encoding="utf-8") as f:
        json.dump(vendor_list, f, ensure_ascii=False)
        
    print("\n=== Build Statistics ===")
    print(f"Total Rows Processed: {total_processed_rows}")
    print(f"Unique Agencies: {len(agencies_set)}")
    print(f"Unique Categories: {len(categories_set)}")
    print(f"Unique Subcategories: {len(subcategories_set)}")
    print(f"Unique Vendors: {len(vendors_set)}")
    print(f"Summary file size: {os.path.getsize(summary_file) / (1024*1024):.2f} MB")
    print(f"Vendors file size: {os.path.getsize(vendors_file) / (1024*1024):.2f} MB")
    print(f"Meta file size: {os.path.getsize(meta_file) / 1024:.2f} KB")

if __name__ == "__main__":
    build_datasets()
